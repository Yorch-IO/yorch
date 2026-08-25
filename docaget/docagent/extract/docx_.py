"""DOCX extraction.

Easier than PDF in the one way that matters: Word carries explicit heading
styles, so the hierarchy does not have to be inferred from numbering and length
guards. A ``Heading 2`` paragraph is a level-2 heading, full stop.

To keep the rest of the pipeline uniform, headings are emitted as standalone
paragraphs prefixed with their numbering when Word supplies none — the chunker's
``heading_level`` only recognises numbered headings, so an unnumbered "Marco
teórico" would otherwise read as body prose.
"""

from __future__ import annotations

import re

from ..chunk import DocRules
from . import Evidence, Extracted, join_paragraphs

_HEADING_STYLE = re.compile(r"^Heading (\d+)$|^Título (\d+)$", re.IGNORECASE)
_ALREADY_NUMBERED = re.compile(r"^\d+(\.\d+)*\.?\s")


def extract(path: str, rules: DocRules) -> Extracted:
    import docx  # python-docx

    doc = docx.Document(path)
    paragraphs: list[str] = []
    # One counter per heading depth, so unnumbered Word headings can be given the
    # numbering the chunker's heading detection expects.
    counters: list[int] = []
    heading_lines: list[str] = []

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        level = _heading_level_from_style(p.style.name if p.style else "")
        if level == 0:
            paragraphs.append(text)
            continue

        if _ALREADY_NUMBERED.match(text):
            numbered = text
        else:
            del counters[level:]
            while len(counters) < level:
                counters.append(0)
            counters[level - 1] += 1
            prefix = ".".join(str(c) for c in counters[:level])
            numbered = f"{prefix}. {text}"
        paragraphs.append(numbered)
        heading_lines.append(numbered)

    # Tables in a DOCX are usually small; serialise each row so its content is at
    # least searchable, with the header row repeated for context.
    for table in doc.tables:
        rows = [[c.text.strip() for c in row.cells] for row in table.rows]
        if not rows:
            continue
        header, body = rows[0], rows[1:]
        for row in body:
            cells = " | ".join(f"{h}: {v}" for h, v in zip(header, row) if v)
            if cells:
                paragraphs.append(cells)

    ev = Evidence(
        source=path,
        extractor="docx_",
        pages=1,
        numbered_lines=heading_lines[:40],
        notes=["headings came from Word styles, not from numbering heuristics"],
    )
    return Extracted(text=join_paragraphs(paragraphs), evidence=ev)


def _heading_level_from_style(style_name: str) -> int:
    if m := _HEADING_STYLE.match(style_name.strip()):
        return int(m.group(1) or m.group(2))
    if style_name.strip().lower() in {"title", "título"}:
        return 1
    return 0
