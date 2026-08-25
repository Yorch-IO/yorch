"""PPTX extraction.

Structured rather than text-like: a deck has no linear byte stream worth slicing,
and its structure is already the breadcrumb the PDF pipeline had to reconstruct —
the slide title. So this extractor emits chunks directly, one per slide, with
``cell_ref`` carrying the slide number as the verifiable anchor.

Speaker notes are included: in lecture decks they routinely hold the substance
that the slide only gestures at.
"""

from __future__ import annotations

from ..chunk import KIND_SLIDE, Chunk, DocRules
from . import Evidence, Extracted


def extract(path: str, rules: DocRules) -> Extracted:
    from pptx import Presentation  # python-pptx

    prs = Presentation(path)
    chunks: list[Chunk] = []
    titles: list[str] = []

    for n, slide in enumerate(prs.slides, start=1):
        title = _title(slide)
        body = _body_text(slide, skip=title)
        notes = _notes(slide)

        parts: list[str] = []
        if body:
            parts.append(body)
        if notes:
            parts.append(f"Notas del orador: {notes}")
        text = "\n\n".join(parts).strip()
        if not text:
            continue

        titles.append(title or f"(diapositiva {n} sin título)")
        chunks.append(
            Chunk(
                index=len(chunks),
                kind=KIND_SLIDE,
                chapter=f"Diapositiva {n}",
                section=title,
                text=text,
                char_from=0,
                char_to=0,
                para_from=n,
                para_to=n,
                cell_ref=f"slide:{n}",
                extra={"slide": n, "has_notes": bool(notes)},
            )
        )

    ev = Evidence(
        source=path,
        extractor="pptx_",
        pages=len(prs.slides.__iter__.__self__._sldIdLst),  # slide count
        numbered_lines=titles[:40],
        notes=["slide titles are used as the breadcrumb; speaker notes included"],
    )
    return Extracted(chunks=chunks, evidence=ev)


def _title(slide) -> str:
    try:
        if slide.shapes.title is not None:
            return (slide.shapes.title.text or "").strip()
    except (AttributeError, ValueError):
        pass
    return ""


def _body_text(slide, skip: str) -> str:
    parts: list[str] = []
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        text = (shape.text_frame.text or "").strip()
        if text and text != skip:
            parts.append(text)
    return "\n".join(parts).strip()


def _notes(slide) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return (frame.text or "").strip() if frame is not None else ""
